"""Read-only mathematical audit. Python 3.11+, openpyxl; optional Node >=22.18.

Run from any directory: python -X utf8 docs/audit-13/verify.py --node /path/to/node
No workbook or application file is modified. JSON is printed to stdout.
"""
import argparse
import ast
from decimal import Decimal as Q, getcontext
import hashlib
import json
import math
from pathlib import Path
import re
import subprocess
import warnings

getcontext().prec = 60
ROOT = Path(__file__).resolve().parents[2]
IDS = 'pi gamma sqrt3 sqrt6 sqrt2 sqrt5 tribonacci brun invPhi phi e eMinus1 eggLW'.split()


def atan_series(x):
    term = x
    result = x
    for n in range(1, 1000):
        term *= -x*x
        add = term / (2*n+1)
        result += add
        if abs(add) < Q('1e-65'):
            return result
    raise AssertionError('atan did not converge')


PI = 16*atan_series(Q(1)/5)-4*atan_series(Q(1)/239)


def atan(x):
    return PI/4 + atan_series((x-1)/(x+1))


def tan_deg(deg):
    x = deg*PI/180
    sine, cosine, st, ct = x, Q(1), x, Q(1)
    for n in range(1, 100):
        st *= -x*x / ((2*n)*(2*n+1))
        ct *= -x*x / ((2*n-1)*(2*n))
        sine += st
        cosine += ct
        if max(abs(st), abs(ct)) < Q('1e-65'):
            return sine/cosine
    raise AssertionError('trigonometric series did not converge')


def bisect(f, a, b):
    fa = f(a)
    assert fa*f(b) <= 0
    for _ in range(210):
        m = (a+b)/2
        fm = f(m)
        if fm == 0:
            return m
        if fa*fm > 0:
            a, fa = m, fm
        else:
            b = m
    return (a+b)/2


PHI = (1+Q(5).sqrt())/2
TRI = bisect(lambda x: x**3-x*x-x-1, Q(1), Q(2))
# NIST DLMF 3.12.E4; fixed reference, not derived from pyramid geometry.
GAMMA = Q('0.577215664901532860606512090082402431042159335939923598805767')
TARGETS = [PI, GAMMA, Q(3).sqrt(), Q(6).sqrt(), Q(2).sqrt(), Q(5).sqrt(),
           TRI, Q('1.902160583104'), 1/PHI, PHI, Q(1).exp(), Q(1).exp()-1, PHI]


def egg(t, z0=Q('7.65')):
    assert 0 < 4*t < z0*z0
    zlo = (z0+(z0*z0-4*t).sqrt())/2
    zhi = (z0+(z0*z0+4*t).sqrt())/2
    # Stationary width from y^2 derivative, independent of engine ternary search.
    zmax = bisect(lambda z: z**3*(z0-z)-t*t, zlo, z0)
    y2 = lambda z: 1/(z*z)-((z-z0)/t)**2
    length = (zhi-zlo)*(1+t*t).sqrt()/t
    width = 2*y2(zmax).sqrt()
    assert abs(y2(zlo)) < Q('1e-55') and abs(y2(zhi)) < Q('1e-55')
    assert zlo < zmax < z0 < zhi and width > 0
    return dict(zlo=zlo, zhi=zhi, zmax=zmax, length=length, width=width, lw=length/width)


def values(t):
    p = (1+t*t).sqrt()
    a = atan(t)
    return [4/t, 4/(t+4*Q(2).sqrt()), (t+4*Q(2).sqrt())/4,
            2+t/(2*Q(2).sqrt()), Q(2).sqrt(), 1+2/p,
            (1+4*Q(2).sqrt())/(p+2), (t*t+2).sqrt(), p/(p+1), p,
            2*a/(PI/2-a), 2*a/(PI/2-a)-1, egg(t)['lw']]


def table(t):
    return [dict(id=i, value=v, reference=c, relative_error=abs(v-c)/c,
                 percent_error=100*abs(v-c)/c, within=abs(v-c)/c <= Q('.001'))
            for i, v, c in zip(IDS, values(t), TARGETS)]


WEIGHTS = [Q(1), Q(1)/3, Q(1)/3, Q(1)/3, Q(0), Q(1)/3,
           Q(1), Q(1), Q(1)/3, Q(1)/3, Q('.5'), Q('.5'), Q(1)]


def independent_scan(step):
    points = []
    for i in range(int(Q('.12')/step)+1):
        angle = Q('51.78')+i*step
        rows = table(tan_deg(angle))
        errors = [r['relative_error'] for r in rows]
        points.append(dict(angle=angle, mean=sum(errors)/13,
                           independent=sum(e*w for e,w in zip(errors,WEIGHTS))/sum(WEIGHTS),
                           rms=(sum(e*e for e in errors)/13).sqrt(),
                           maximum=max(errors), matches=sum(r['within'] for r in rows),
                           classical_matches=sum(r['within'] for r in rows[:12])))
    nearby = [p for p in points if abs(p['angle']-atan(Q(14)/11)*180/PI) <= Q('.02')]
    return dict(step=step, count=len(points),
                minima={k:min(points,key=lambda p:p[k])['angle']
                        for k in ['mean','independent','rms','maximum']},
                all13=[p['angle'] for p in points if p['matches']==13],
                near_11_7=dict(total=len(nearby), all13=sum(p['matches']==13 for p in nearby),
                              all12=sum(p['classical_matches']==12 for p in nearby)))


def workbook_audit(path, expected):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', UserWarning)
        w = load_workbook(path, data_only=False)
        cached = load_workbook(path, data_only=True)
    sheets = {s.title[:2]: s.title for s in w}
    records = {}
    active = set()

    def cell(sheet, addr):
        addr = addr.replace('$', '')
        key = sheet+'!'+addr
        raw = w[sheet][addr].value
        records[key] = dict(formula_or_input=raw, cached=cached[sheet][addr].value)
        if not isinstance(raw, str) or not raw.startswith('='):
            assert isinstance(raw, (float, int)), (key, raw)
            return float(raw)
        assert key not in active, key
        active.add(key)
        expr = raw[1:].replace('^', '**')
        expr = re.sub(r"'([^']+)'!(\$?[A-Z]+\$?\d+)",
                      lambda m: repr(cell(m[1], m[2])), expr)
        expr = re.sub(r'(?<![A-Za-z0-9_])\$?[A-Z]+\$?\d+',
                      lambda m: repr(cell(sheet, m[0])), expr)
        funcs = dict(SQRT=math.sqrt, ATAN=math.atan, DEGREES=math.degrees,
                     RADIANS=math.radians, TAN=math.tan, SIN=math.sin,
                     COS=math.cos, ABS=abs, PI=lambda: math.pi, EXP=math.exp)
        tree = ast.parse(expr, mode='eval')
        for n in ast.walk(tree):
            assert isinstance(n, (ast.Expression, ast.BinOp, ast.UnaryOp, ast.Add,
                                 ast.Sub, ast.Mult, ast.Div, ast.Pow, ast.USub,
                                 ast.UAdd, ast.Constant, ast.Call, ast.Name, ast.Load)), (key, expr)
            if isinstance(n, ast.Name):
                assert n.id in funcs
            if isinstance(n, ast.Call):
                assert isinstance(n.func, ast.Name) and n.func.id in funcs
        result = eval(compile(tree, '<restricted workbook arithmetic>', 'eval'),
                      {'__builtins__': {}}, funcs)
        active.remove(key)
        return result

    deltas = []
    for row, t in [(7, Q(14)/11), (12, expected['golden_t'])]:
        vs = values(t)
        for index in range(12):
            addr = get_column_letter(4+index)+str(row)
            val = cell(sheets['03'], addr)
            delta = abs(val-float(vs[index]))
            assert delta < 5e-13, (path.name, addr, val, vs[index])
            deltas.append(delta)
            target = cell(sheets['02'], 'D'+str(6+index))
            err = cell(sheets['03'], get_column_letter(16+index)+str(row))
            assert abs(err-abs(val-target)/target) < 1e-14
            assert abs(target-float(TARGETS[index])) < 5e-15
    egg_sheet = sheets['11']
    lw = cell(egg_sheet, 'R25')
    assert abs(lw-float(PHI)) < 1e-12
    zmax = expected['golden_egg']['zmax']
    t = expected['golden_t']
    vmax = (Q('7.65')-zmax)*(1+t*t).sqrt()/t
    frozen_v = cell(egg_sheet, 'R27')
    # Width is insensitive at its maximum; the stored v is not an exact solver.
    frozen_v_delta = abs(Q(str(frozen_v))-vmax)
    score_sheet = sheets['04']
    for addr in ['M7', 'N7', 'O7']:
        records[score_sheet+'!'+addr] = dict(formula_or_input=w[score_sheet][addr].value,
                                             cached=cached[score_sheet][addr].value)
    for addr in ['B7','B8','B9','B10','B11','B12','B13','E6','E7','E8','E9']:
        cell('00_Start', addr)
    w.close()
    cached.close()
    return dict(file=path.name, sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                max_classical_value_delta=max(deltas), frozen_v_delta=frozen_v_delta,
                sheets=sheets, evidence=records)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--node', help='Node >=22.18 supporting direct .ts import')
    args = parser.parse_args()
    t = Q(14)/11
    golden_t = bisect(lambda x: egg(x)['lw']-PHI, Q('1.26'), Q('1.28'))
    base = table(t)
    golden = table(golden_t)
    assert sum(r['within'] for r in base) == 12
    assert sum(r['within'] for r in golden) == 10
    assert abs(TRI**3-TRI**2-TRI-1) < Q('1e-55')
    # Geometric dependencies at several different slopes, not only 11:7.
    for slope in map(Q, ['0.5', '1', '1.2727272727272727', '2']):
        v = values(slope)
        expected = [(v[1]*v[2], Q(1)), (v[3], v[2]*v[4]),
                    (v[5], 1+2/v[9]), (v[6], (1+4*v[4])/(v[9]+2)),
                    (v[7]**2, v[9]**2+1), (v[8], v[9]/(v[9]+1)),
                    (v[11], v[10]-1)]
        assert all(abs(a-b) < Q('1e-55') for a,b in expected)
    # Cartesian Pythagoras/scale checks against reduced one-parameter formulas.
    for b in map(Q, ['1', '11', '72.0225', '440']):
        h, a = b*7/11, b/2
        s, e, d = (h*h+a*a).sqrt(), (h*h+2*a*a).sqrt(), b*Q(2).sqrt()
        direct = [2*b/h, 2*b/(h+2*d), (h+2*d)/(2*b), (h+2*d)/d,
                  d/b, (s+b)/s, (a+2*d)/(s+b), e/a, s/(s+a), s/a]
        assert all(abs(x-r['value']) < Q('1e-55') for x,r in zip(direct,base))
    report = dict(precision=60, angle_11_7=atan(t)*180/PI,
                  golden_t=golden_t, golden_angle=atan(golden_t)*180/PI,
                  eleven_seven=base, golden=golden,
                  egg_11_7=egg(t), golden_egg=egg(golden_t))
    report['independent_scans'] = [independent_scan(Q(s)) for s in ['.0005','.001']]
    report['workbooks'] = [workbook_audit(p, report) for p in sorted((ROOT/'data').glob('*.xlsx'))]
    assert len(report['workbooks']) == 2
    if args.node:
        # Import production formulas only for comparison with the independent audit.
        code = "import * as m from " + json.dumps((ROOT/'src/lib/pyramid/engine.ts').as_uri()) + ";"
        code += "console.log(JSON.stringify({models:['elevenSeven','goldenEgg'].map(id=>m.snapshotFor(m.MODELS.find(x=>x.id===id),11/7)),scans:[0.0005,0.001].map(step=>{const p=m.scanAngles(undefined,undefined,step);return {step,count:p.length,minima:m.scanMinima(p),all13:p.filter(x=>x.matches===13).map(x=>x.angle)}})}));"
        actual = json.loads(subprocess.check_output([args.node, '--input-type=module', '-e', code], text=True))
        for got, ref in zip(actual['scans'], report['independent_scans']):
            assert got['count'] == ref['count']
            assert len(got['all13']) == len(ref['all13'])
            assert all(abs(Q(str(a))-b)<Q('1e-12') for a,b in zip(got['all13'],ref['all13']))
            for name in ['mean','independent','minimax','rmsAngle']:
                target_name = {'minimax': 'maximum', 'rmsAngle': 'rms'}.get(name, name)
                assert abs(Q(str(got['minima'][name]))-ref['minima'][target_name]) < Q('1e-12')
        max_delta = 0
        for snapshot, rows in zip(actual['models'], [base,golden]):
            assert [r['id'] for r in snapshot['results']] == IDS
            for got, ref in zip(snapshot['results'], rows):
                delta = abs(Q(str(got['value']))-ref['value'])
                max_delta = max(max_delta, delta)
                # Golden preset angle is rounded to 9 decimals in production.
                assert delta < Q('2e-11'), (got['id'], delta)
                assert got['within'] == ref['within']
        report['engine_comparison'] = dict(max_value_delta=max_delta, output=actual)
    print(json.dumps(report, ensure_ascii=False, indent=2, default=str))


if __name__ == '__main__':
    main()
